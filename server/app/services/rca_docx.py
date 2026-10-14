"""RCA brief as a Word document."""
from io import BytesIO

from docx import Document as DocxDocument
from docx.shared import Pt


def build_rca_docx(rca: dict) -> bytes:
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    style.font.name = "Calibri"

    doc.add_heading(f"Root Cause Analysis — {rca['assetTag']}", level=0)
    strength = (rca.get("evidenceStrength") or "low").capitalize()
    origin = f"drafted by {rca['model']}" if rca.get("mode") == "llm" else "extractive summary of the evidence"
    doc.add_paragraph(f"Evidence strength: {strength} ({origin})")
    if rca.get("notice"):
        doc.add_paragraph(rca["notice"])
    for warning in rca.get("warnings") or []:
        doc.add_paragraph(warning, style="List Bullet")

    doc.add_heading("Observed Problem", level=1)
    doc.add_paragraph(rca["observedProblem"])

    sections = [
        ("Probable Causes", rca["probableCauses"], "List Bullet"),
        ("Recorded Corrective Actions", rca["recordedActions"], "List Bullet"),
        ("Relevant Measurements", rca["measurements"], "List Bullet"),
    ]
    for title, items, style_name in sections:
        if items:
            doc.add_heading(title, level=1)
            for item in items:
                doc.add_paragraph(item, style=style_name)

    if rca["eventDates"]:
        doc.add_heading("Event Dates", level=1)
        doc.add_paragraph(", ".join(rca["eventDates"]))

    for title, key in (("Recommended Investigation", "recommendedInvestigation"), ("Preventive Actions", "preventiveActions")):
        doc.add_heading(title, level=1)
        for item in rca[key]:
            doc.add_paragraph(item, style="List Number")

    if rca.get("citations"):
        doc.add_heading("Evidence", level=1)
        for c in rca["citations"]:
            paragraph = doc.add_paragraph(style="List Bullet")
            label = f"[{c['label']}] " if c.get("label") else ""
            paragraph.add_run(f"{label}{c['source']} (p. {c['page']}): ").bold = True
            paragraph.add_run(c["excerpt"])

    doc.add_paragraph("")
    doc.add_paragraph().add_run(rca["disclaimer"]).italic = True

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
