"""Text extraction from industrial documents: PDF (with OCR fallback), DOCX, CSV, XLSX, PPTX,
HTML, EML and images. Entity extraction lives in entities.py."""
import csv
from pathlib import Path

def extract_text(path: Path) -> tuple[str, list[dict]]:
    """Return (full_text, pages). Raises ValueError when nothing readable can be extracted,
    so the upload is marked FAILED instead of indexing an error message as content."""
    text, pages = _extract(path)
    if not text.strip():
        raise ValueError(
            "No readable text found. Scanned PDFs and images need Tesseract OCR (set TESSERACT_CMD)."
        )
    return text, pages


def _extract(path: Path) -> tuple[str, list[dict]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf(path)
    if suffix == ".docx":
        return _docx(path)
    if suffix == ".csv":
        return _csv(path)
    if suffix in {".jpg", ".jpeg", ".png", ".tiff", ".bmp"}:
        return _image(path)
    if suffix in {".xlsx", ".xls"}:
        return _xlsx(path)
    if suffix == ".pptx":
        return _pptx(path)
    if suffix in {".html", ".htm"}:
        return _html(path)
    if suffix in {".eml", ".msg"}:
        return _eml(path)
    text = path.read_text(encoding="utf-8", errors="ignore")
    return text, [{"page": 1, "text": text}]


def _pdf(path: Path) -> tuple[str, list[dict]]:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            pages.append({"page": index, "text": text})
        full_text = "\n\n".join(item["text"] for item in pages)
        if len(full_text.strip()) < 30:
            ocr_text = _ocr_pdf(path)
            if ocr_text:
                pages = [{"page": i + 1, "text": t} for i, t in enumerate(ocr_text)]
                full_text = "\n\n".join(ocr_text)
        return full_text, pages
    except Exception as e:
        raise ValueError(f"PDF extraction failed: {e}") from e


def _ocr_pdf(path: Path) -> list[str]:
    try:
        import pytesseract
        from pdf2image import convert_from_path
        from ..core.config import get_settings
        if get_settings().tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = get_settings().tesseract_cmd
        return [pytesseract.image_to_string(img) for img in convert_from_path(str(path), dpi=180)]
    except Exception:
        return []


def _docx(path: Path) -> tuple[str, list[dict]]:
    try:
        from docx import Document
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return text, [{"page": 1, "text": text}]
    except Exception as e:
        raise ValueError(f"DOCX extraction failed: {e}") from e


def _csv(path: Path) -> tuple[str, list[dict]]:
    rows = []
    with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as f:
        for row in csv.reader(f):
            rows.append(" | ".join(row))
    text = "\n".join(rows)
    return text, [{"page": 1, "text": text}]


def _image(path: Path) -> tuple[str, list[dict]]:
    try:
        import pytesseract
        from PIL import Image
        from ..core.config import get_settings
        if get_settings().tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = get_settings().tesseract_cmd
        text = pytesseract.image_to_string(Image.open(str(path)))
        return text, [{"page": 1, "text": text}]
    except Exception as e:
        raise ValueError(f"Image OCR failed: {e}") from e


def _xlsx(path: Path) -> tuple[str, list[dict]]:
    try:
        import pandas as pd
        df = pd.read_excel(str(path))
        rows = []
        for _, row in df.iterrows():
            rows.append(" | ".join(f"{col}: {val}" for col, val in row.items() if pd.notna(val)))
        text = "\n".join(rows)
        return text, [{"page": 1, "text": text}]
    except Exception as e:
        raise ValueError(f"Excel parsing failed: {e}") from e


def _pptx(path: Path) -> tuple[str, list[dict]]:
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        pages = []
        for idx, slide in enumerate(prs.slides, start=1):
            slide_text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        if p.text.strip():
                            slide_text.append(p.text.strip())
            pages.append({"page": idx, "text": "\n".join(slide_text)})
        full_text = "\n\n".join(item["text"] for item in pages)
        return full_text, pages
    except Exception as e:
        raise ValueError(f"PPTX parsing failed: {e}") from e


def _html(path: Path) -> tuple[str, list[dict]]:
    try:
        from bs4 import BeautifulSoup
        text = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser").get_text(separator="\n")
        return text, [{"page": 1, "text": text}]
    except Exception:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return text, [{"page": 1, "text": text}]


def _eml(path: Path) -> tuple[str, list[dict]]:
    try:
        import email
        from email import policy
        msg = email.message_from_string(path.read_text(encoding="utf-8", errors="ignore"), policy=policy.default)
        body = ""
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    body += part.get_content()
        else:
            body = msg.get_content()
        text = f"Subject: {msg.get('subject', '')}\nFrom: {msg.get('from', '')}\nDate: {msg.get('date', '')}\n\n{body}"
        return text, [{"page": 1, "text": text}]
    except Exception:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return text, [{"page": 1, "text": text}]
